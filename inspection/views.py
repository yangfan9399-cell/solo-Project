from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse, JsonResponse, Http404
from django.views.decorators.http import require_http_methods, require_GET
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q, Count, Sum, Avg, Min, Max, F
from django.db.models.functions import TruncDate, TruncWeek, ExtractWeekDay
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.template.loader import render_to_string
from django.views.decorators.csrf import csrf_exempt
from datetime import datetime, date, timedelta
from decimal import Decimal
import csv
import io
import json

from .models import CrystallizationPool, InspectionRecord, RecordVersion, AnomalyAlert
from .forms import (
    InspectionRecordForm, InspectionFilterForm, CrystallizationPoolForm, AnomalyAlertForm
)


def _get_date_range(request):
    today = timezone.localdate()
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    if not date_from:
        date_from = (today - timedelta(days=30)).isoformat()
    if not date_to:
        date_to = today.isoformat()
    return parse_date(date_from) or (today - timedelta(days=30)), parse_date(date_to) or today


def _apply_filters(queryset, request):
    form = InspectionFilterForm(request.GET or None)
    qs = queryset
    if request.GET:
        keyword = request.GET.get('keyword', '').strip()
        if keyword:
            qs = qs.filter(
                Q(pool__pool_code__icontains=keyword) |
                Q(pool__pool_name__icontains=keyword) |
                Q(batch_no__icontains=keyword) |
                Q(inspector__icontains=keyword) |
                Q(anomaly_description__icontains=keyword)
            )
        pool_group = request.GET.get('pool_group')
        if pool_group:
            qs = qs.filter(pool__pool_group=pool_group)
        pool_id = request.GET.get('pool_id')
        if pool_id and pool_id.isdigit():
            qs = qs.filter(pool_id=int(pool_id))
        date_from = request.GET.get('date_from')
        if date_from:
            d = parse_date(date_from)
            if d:
                qs = qs.filter(inspection_date__gte=d)
        date_to = request.GET.get('date_to')
        if date_to:
            d = parse_date(date_to)
            if d:
                qs = qs.filter(inspection_date__lte=d)
        surface = request.GET.get('surface_status')
        if surface:
            qs = qs.filter(surface_status=surface)
        weather = request.GET.get('weather_type')
        if weather:
            qs = qs.filter(weather_type=weather)
        if request.GET.get('anomaly_only') == 'on' or request.GET.get('anomaly_only') == 'true':
            qs = qs.filter(has_anomaly=True)
        cmin = request.GET.get('concentration_min')
        if cmin:
            try:
                qs = qs.filter(brine_concentration__gte=Decimal(cmin))
            except Exception:
                pass
        cmax = request.GET.get('concentration_max')
        if cmax:
            try:
                qs = qs.filter(brine_concentration__lte=Decimal(cmax))
            except Exception:
                pass
    return qs, form


def dashboard(request):
    today = timezone.localdate()
    last_30 = today - timedelta(days=30)
    last_7 = today - timedelta(days=7)

    pools_total = CrystallizationPool.objects.count()
    pools_active = CrystallizationPool.objects.filter(status='active').count()
    records_30 = InspectionRecord.objects.filter(inspection_date__gte=last_30)

    total_records = records_30.count()
    avg_conc = records_30.aggregate(a=Avg('brine_concentration'))['a'] or Decimal('0')
    total_yield = records_30.aggregate(s=Sum('salt_yield'))['s'] or Decimal('0')
    anomaly_count = AnomalyAlert.objects.filter(status__in=('open', 'processing')).count()
    anomaly_critical = AnomalyAlert.objects.filter(severity__in=('high', 'critical'), status__in=('open',)).count()

    latest_records = InspectionRecord.objects.select_related('pool')[:20]

    groups = ['A', 'B', 'C', 'D']
    group_stats = []
    for g in groups:
        g_pools = CrystallizationPool.objects.filter(pool_group=g)
        g_ids = g_pools.values_list('id', flat=True)
        g_records = records_30.filter(pool_id__in=g_ids)
        group_stats.append({
            'group': g,
            'pool_count': g_pools.count(),
            'active_count': g_pools.filter(status='active').count(),
            'records': g_records.count(),
            'avg_conc': (g_records.aggregate(a=Avg('brine_concentration'))['a'] or Decimal('0')),
            'total_yield': (g_records.aggregate(s=Sum('salt_yield'))['s'] or Decimal('0')),
        })

    pools = CrystallizationPool.objects.all().order_by('pool_code')
    pool_status_map = {}
    for p in pools:
        latest = p.inspections.first()
        pool_status_map[p.id] = {
            'id': p.id,
            'code': p.pool_code,
            'name': p.pool_name,
            'group': p.pool_group,
            'x': p.position_x,
            'y': p.position_y,
            'area': float(p.area),
            'status': p.status,
            'latest_conc': float(latest.brine_concentration) if latest else None,
            'latest_date': latest.inspection_date.isoformat() if latest else None,
            'latest_surface': latest.surface_status if latest else None,
            'has_anomaly': latest.has_anomaly if latest else False,
            'yield_today': float(latest.salt_yield) if latest and latest.inspection_date == today else 0,
        }

    recent_alerts = AnomalyAlert.objects.select_related('pool').filter(
        status__in=('open', 'processing')
    ).order_by('-severity', '-detected_at')[:10]

    context = {
        'today': today,
        'pools_total': pools_total,
        'pools_active': pools_active,
        'total_records': total_records,
        'avg_conc': round(float(avg_conc), 2),
        'total_yield': round(float(total_yield), 3),
        'anomaly_count': anomaly_count,
        'anomaly_critical': anomaly_critical,
        'group_stats': group_stats,
        'pool_status_json': json.dumps(list(pool_status_map.values())),
        'latest_records': latest_records,
        'recent_alerts': recent_alerts,
        'last_30': last_30,
        'last_7': last_7,
    }
    return render(request, 'inspection/dashboard.html', context)


def record_list(request):
    qs = InspectionRecord.objects.select_related('pool').all()
    qs, form = _apply_filters(qs, request)

    paginator = Paginator(qs, 25)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except Exception:
        page_obj = paginator.page(1)

    stats = qs.aggregate(
        total=Count('id'),
        avg_conc=Avg('brine_concentration'),
        sum_yield=Sum('salt_yield'),
        anomaly_count=Count('id', filter=Q(has_anomaly=True)),
    )

    date_from, date_to = _get_date_range(request)
    context = {
        'form': form,
        'page_obj': page_obj,
        'stats': stats,
        'date_from': date_from,
        'date_to': date_to,
        'pool_groups': CrystallizationPool.POOL_GROUP,
        'surface_choices': InspectionRecord.SURFACE_STATUS,
        'weather_choices': InspectionRecord.WEATHER_TYPE,
        'query_string': request.GET.urlencode(),
    }
    return render(request, 'inspection/record_list.html', context)


def record_detail(request, pk):
    record = get_object_or_404(InspectionRecord.objects.select_related('pool'), pk=pk)
    versions = record.versions.all().order_by('-version_no')
    alerts = AnomalyAlert.objects.filter(record=record).order_by('-severity')

    pool_qs = InspectionRecord.objects.filter(pool=record.pool, inspection_date__lte=record.inspection_date).order_by('-inspection_date')[:10]
    related_records = list(pool_qs)

    concentration_val = float(record.brine_concentration)
    conc_status = 'normal'
    if concentration_val < 22 or concentration_val > 31:
        conc_status = 'danger'
    elif concentration_val < 24 or concentration_val > 30:
        conc_status = 'warning'

    surface_ok = record.surface_status in ('normal', 'crystallized')

    context = {
        'record': record,
        'versions': versions,
        'alerts': alerts,
        'related_records': related_records,
        'conc_status': conc_status,
        'surface_ok': surface_ok,
    }
    return render(request, 'inspection/record_detail.html', context)


def record_create(request):
    if request.method == 'POST':
        form = InspectionRecordForm(request.POST, request.FILES)
        if form.is_valid():
            record = form.save(commit=False)
            record.created_by = request.user if request.user.is_authenticated else None
            record.save()
            RecordVersion.objects.create(
                record=record,
                version_no=1,
                action='create',
                change_summary='手动创建巡检记录',
                snapshot={
                    'concentration': float(record.brine_concentration),
                    'surface': record.surface_status,
                    'yield': float(record.salt_yield),
                },
            )
            messages.success(request, f'已创建巡检记录 #{record.id}')
            return redirect('inspection:record_detail', pk=record.id)
    else:
        form = InspectionRecordForm(initial={
            'inspection_date': timezone.localdate(),
            'inspection_time': timezone.localtime().time(),
        })
    context = {'form': form, 'is_edit': False}
    return render(request, 'inspection/record_form.html', context)


def record_edit(request, pk):
    record = get_object_or_404(InspectionRecord, pk=pk)
    old_snapshot = {
        'concentration': float(record.brine_concentration),
        'surface': record.surface_status,
        'yield': float(record.salt_yield),
        'weather': record.weather_type,
        'depth': float(record.brine_depth_cm),
        'crystal': float(record.crystal_thickness_mm),
    }
    if request.method == 'POST':
        form = InspectionRecordForm(request.POST, request.FILES, instance=record)
        if form.is_valid():
            new_record = form.save(commit=False)
            new_version = record.version + 1
            new_record.version = new_version
            new_record.save()

            changes = []
            new_conc = float(new_record.brine_concentration)
            new_surf = new_record.surface_status
            new_yield = float(new_record.salt_yield)
            if new_conc != old_snapshot['concentration']:
                changes.append(f'浓度{old_snapshot["concentration"]}→{new_conc}°Bé')
            if new_surf != old_snapshot['surface']:
                changes.append(f'池面{dict(InspectionRecord.SURFACE_STATUS).get(old_snapshot["surface"])}→{new_record.get_surface_status_display()}')
            if abs(new_yield - old_snapshot['yield']) > 0.001:
                changes.append(f'产量{old_snapshot["yield"]}→{new_yield}吨')
            if not changes:
                changes.append('修改元数据/备注信息')

            RecordVersion.objects.create(
                record=new_record,
                version_no=new_version,
                action='update',
                change_summary='；'.join(changes),
                snapshot={
                    'old': old_snapshot,
                    'new': {
                        'concentration': new_conc,
                        'surface': new_surf,
                        'yield': new_yield,
                        'weather': new_record.weather_type,
                    },
                },
                changed_by=request.user if request.user.is_authenticated else None,
            )
            messages.success(request, f'记录已更新至版本 v{new_version}')
            return redirect('inspection:record_detail', pk=record.id)
    else:
        form = InspectionRecordForm(instance=record)
    context = {'form': form, 'record': record, 'is_edit': True}
    return render(request, 'inspection/record_form.html', context)


def record_delete(request, pk):
    record = get_object_or_404(InspectionRecord, pk=pk)
    pool_code = record.pool.pool_code
    insp_date = record.inspection_date
    record.delete()
    messages.info(request, f'已删除 {pool_code} 在 {insp_date} 的巡检记录')
    return redirect('inspection:record_list')


def record_versions(request, pk):
    record = get_object_or_404(InspectionRecord.objects.select_related('pool'), pk=pk)
    versions = record.versions.all().order_by('-version_no')
    versions_list = []
    prev = None
    for v in versions:
        item = {
            'version': v,
            'diff': [],
        }
        if prev and v.snapshot and prev.snapshot:
            old = v.snapshot
            new = prev.snapshot
            if isinstance(old.get('new'), dict):
                old = old['new']
            if isinstance(new.get('old'), dict):
                new = new['old']
            for key in ['concentration', 'surface', 'yield']:
                if key in old and key in new and old[key] != new[key]:
                    item['diff'].append((key, old[key], new[key]))
        versions_list.append(item)
        prev = v

    context = {
        'record': record,
        'versions_list': versions_list,
    }
    return render(request, 'inspection/record_versions.html', context)


def record_revert(request, pk, version_no):
    record = get_object_or_404(InspectionRecord, pk=pk)
    target_version = get_object_or_404(RecordVersion, record=record, version_no=version_no)
    snap = target_version.snapshot or {}
    data = snap.get('new') if isinstance(snap.get('new'), dict) else snap

    old_snap = {
        'concentration': float(record.brine_concentration),
        'surface': record.surface_status,
        'yield': float(record.salt_yield),
    }

    if 'concentration' in data:
        record.brine_concentration = data['concentration']
    if 'surface' in data:
        record.surface_status = data['surface']
    if 'yield' in data:
        record.salt_yield = data['yield']
    record.version = record.version + 1
    record.save()

    RecordVersion.objects.create(
        record=record,
        version_no=record.version,
        action='revert',
        change_summary=f'回滚至版本 v{version_no}',
        snapshot={
            'revert_to': version_no,
            'old': old_snap,
            'new': data,
        },
        changed_by=request.user if request.user.is_authenticated else None,
    )
    messages.success(request, f'已回滚记录至 v{version_no}，当前版本 v{record.version}')
    return redirect('inspection:record_versions', pk=pk)


def batch_list(request):
    today = timezone.localdate()
    last_90 = today - timedelta(days=90)

    records_with_batch = InspectionRecord.objects.filter(
        batch_no__isnull=False,
    ).exclude(batch_no='').filter(inspection_date__gte=last_90)

    batches = records_with_batch.values('batch_no', 'pool__pool_group').annotate(
        record_count=Count('id'),
        first_date=Min('inspection_date'),
        last_date=Max('inspection_date'),
        total_yield=Sum('salt_yield'),
        avg_conc=Avg('brine_concentration'),
        pools=Count('pool', distinct=True),
        anomaly_count=Count('id', filter=Q(has_anomaly=True)),
    ).order_by('-first_date')

    paginator = Paginator(batches, 30)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except Exception:
        page_obj = paginator.page(1)

    context = {'page_obj': page_obj}
    return render(request, 'inspection/batch_list.html', context)


def alert_list(request):
    qs = AnomalyAlert.objects.select_related('pool', 'record').all()
    status = request.GET.get('status')
    severity = request.GET.get('severity')
    atype = request.GET.get('anomaly_type')
    if status:
        qs = qs.filter(status=status)
    if severity:
        qs = qs.filter(severity=severity)
    if atype:
        qs = qs.filter(anomaly_type=atype)

    paginator = Paginator(qs.order_by('-created_at'), 30)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except Exception:
        page_obj = paginator.page(1)

    stats = {
        'total': qs.count(),
        'open': qs.filter(status='open').count(),
        'processing': qs.filter(status='processing').count(),
        'critical': qs.filter(severity='critical').count(),
    }

    context = {
        'page_obj': page_obj,
        'stats': stats,
        'filter_status': status,
        'filter_severity': severity,
        'filter_type': atype,
    }
    return render(request, 'inspection/alert_list.html', context)


def alert_resolve(request, pk):
    alert = get_object_or_404(AnomalyAlert, pk=pk)
    alert.status = 'resolved'
    alert.resolved_at = timezone.now()
    alert.handler = alert.handler or (request.user.get_full_name() if request.user.is_authenticated else '系统')
    alert.resolution = alert.resolution or '已按标准流程处理并确认恢复正常'
    alert.save()
    messages.success(request, f'异常告警 #{alert.id} 已标记为已解决')
    return redirect(request.META.get('HTTP_REFERER') or 'inspection:alert_list')


def pool_list(request):
    pools = CrystallizationPool.objects.all().annotate(
        record_count=Count('inspections'),
        latest_conc=Max('inspections__brine_concentration'),
        total_yield=Sum('inspections__salt_yield'),
    ).order_by('pool_code')
    context = {'pools': pools}
    return render(request, 'inspection/pool_list.html', context)


def pool_detail(request, pk):
    pool = get_object_or_404(CrystallizationPool, pk=pk)
    last_30 = timezone.localdate() - timedelta(days=30)
    records = pool.inspections.filter(inspection_date__gte=last_30).order_by('-inspection_date')
    stats = records.aggregate(
        count=Count('id'),
        avg_conc=Avg('brine_concentration'),
        min_conc=Min('brine_concentration'),
        max_conc=Max('brine_concentration'),
        total_yield=Sum('salt_yield'),
        anomaly_count=Count('id', filter=Q(has_anomaly=True)),
    )
    alerts = AnomalyAlert.objects.filter(pool=pool).order_by('-detected_at')[:10]
    context = {
        'pool': pool,
        'records': records[:30],
        'stats': stats,
        'alerts': alerts,
    }
    return render(request, 'inspection/pool_detail.html', context)


@require_GET
def api_pool_status(request):
    pools = CrystallizationPool.objects.all()
    data = []
    for p in pools:
        latest = p.inspections.first()
        data.append({
            'id': p.id,
            'code': p.pool_code,
            'name': p.pool_name,
            'group': p.pool_group,
            'x': p.position_x,
            'y': p.position_y,
            'area': float(p.area),
            'status': p.status,
            'latest_conc': float(latest.brine_concentration) if latest else None,
            'latest_date': latest.inspection_date.isoformat() if latest else None,
            'surface': latest.surface_status if latest else None,
            'has_anomaly': latest.has_anomaly if latest else False,
        })
    return JsonResponse({'pools': data})


@require_GET
def api_trend_data(request):
    days = int(request.GET.get('days', 30))
    pool_id = request.GET.get('pool_id')
    today = timezone.localdate()
    start = today - timedelta(days=days)
    qs = InspectionRecord.objects.filter(inspection_date__gte=start)
    if pool_id and pool_id.isdigit():
        qs = qs.filter(pool_id=int(pool_id))

    daily = qs.values('inspection_date').annotate(
        avg_conc=Avg('brine_concentration'),
        min_conc=Min('brine_concentration'),
        max_conc=Max('brine_concentration'),
        daily_yield=Sum('salt_yield'),
        records=Count('id'),
    ).order_by('inspection_date')

    labels = []
    conc_avg = []
    yield_series = []
    conc_min = []
    conc_max = []
    for d in daily:
        labels.append(d['inspection_date'].strftime('%m-%d'))
        conc_avg.append(round(float(d['avg_conc'] or 0), 2))
        conc_min.append(round(float(d['min_conc'] or 0), 2))
        conc_max.append(round(float(d['max_conc'] or 0), 2))
        yield_series.append(round(float(d['daily_yield'] or 0), 3))

    return JsonResponse({
        'labels': labels,
        'conc_avg': conc_avg,
        'conc_min': conc_min,
        'conc_max': conc_max,
        'daily_yield': yield_series,
    })


def _write_csv_response(rows, headers, filename):
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow(headers)
    writer.writerows(rows)
    return response


def export_csv(request):
    qs = InspectionRecord.objects.select_related('pool').all()
    qs, _ = _apply_filters(qs, request)
    qs = qs[:5000]
    rows = []
    for r in qs:
        rows.append([
            r.id,
            r.inspection_date.isoformat(),
            r.inspection_time.strftime('%H:%M'),
            r.pool.pool_code,
            r.pool.pool_name,
            r.pool.pool_group,
            r.batch_no,
            r.inspector,
            float(r.brine_concentration),
            float(r.brine_depth_cm),
            float(r.crystal_thickness_mm),
            r.get_surface_status_display(),
            r.get_weather_type_display(),
            float(r.temperature),
            float(r.humidity),
            r.get_wind_level_display(),
            r.wind_direction,
            float(r.salt_yield),
            float(r.cumulative_yield),
            float(r.ph_value) if r.ph_value else '',
            r.get_impurity_level_display() if r.impurity_level else '',
            '是' if r.has_anomaly else '否',
            r.anomaly_description,
            r.remarks,
            r.version,
        ])
    headers = ['记录ID', '巡检日期', '巡检时间', '池号', '池名', '池组', '批次号', '巡检员',
               '浓度(°Bé)', '卤水深度(cm)', '结晶厚度(mm)', '池面状态', '天气', '气温(℃)',
               '湿度(%)', '风力', '风向', '本次收盐(吨)', '累计产量(吨)', 'pH值', '杂质等级',
               '是否异常', '异常描述', '备注', '版本号']
    return _write_csv_response(rows, headers, f'巡检台账导出_{timezone.localdate().isoformat()}.csv')


def export_excel(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        messages.error(request, '请先安装 openpyxl：pip install openpyxl')
        return redirect('inspection:record_list')

    qs = InspectionRecord.objects.select_related('pool').all()
    qs, _ = _apply_filters(qs, request)
    qs = qs[:5000]

    wb = Workbook()
    ws = wb.active
    ws.title = '巡检台账'

    headers = ['记录ID', '巡检日期', '巡检时间', '池号', '池名', '池组', '批次号', '巡检员',
               '浓度(°Bé)', '卤水深度(cm)', '结晶厚度(mm)', '池面状态', '天气', '气温(℃)',
               '湿度(%)', '风力', '风向', '本次收盐(吨)', '累计产量(吨)', 'pH值', '杂质等级',
               '是否异常', '异常描述', '备注', '版本号']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(border_style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    red_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    for row_idx, r in enumerate(qs, 2):
        data = [
            r.id,
            r.inspection_date.isoformat(),
            r.inspection_time.strftime('%H:%M'),
            r.pool.pool_code,
            r.pool.pool_name,
            r.pool.pool_group,
            r.batch_no,
            r.inspector,
            float(r.brine_concentration),
            float(r.brine_depth_cm),
            float(r.crystal_thickness_mm),
            r.get_surface_status_display(),
            r.get_weather_type_display(),
            float(r.temperature),
            float(r.humidity),
            r.get_wind_level_display(),
            r.wind_direction,
            float(r.salt_yield),
            float(r.cumulative_yield),
            float(r.ph_value) if r.ph_value else None,
            r.get_impurity_level_display() if r.impurity_level else '',
            '是' if r.has_anomaly else '否',
            r.anomaly_description,
            r.remarks,
            r.version,
        ]
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.alignment = center
            c.border = border
            if r.has_anomaly:
                c.fill = red_fill

    widths = [8, 12, 10, 10, 12, 6, 20, 10, 10, 12, 14, 10, 8, 8, 8, 8, 8, 14, 14, 8, 10, 10, 30, 20, 8]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else 'A' + chr(64 + col - 26)].width = w

    today_str = timezone.localdate().isoformat()
    summary_ws = wb.create_sheet(f'摘要-{today_str}')
    date_from, date_to = _get_date_range(request)
    summary_stats = InspectionRecord.objects.filter(
        inspection_date__gte=date_from, inspection_date__lte=date_to
    ).aggregate(
        records=Count('id'),
        pools=Count('pool', distinct=True),
        avg_conc=Avg('brine_concentration'),
        min_conc=Min('brine_concentration'),
        max_conc=Max('brine_concentration'),
        total_yield=Sum('salt_yield'),
        anomalies=Count('id', filter=Q(has_anomaly=True)),
        batches=Count('batch_no', distinct=True, filter=~Q(batch_no='')),
    )
    summary_data = [
        ['海盐晒场结晶池巡检台账 - 导出摘要'],
        ['导出日期', today_str],
        ['统计区间', f'{date_from.isoformat()} ~ {date_to.isoformat()}'],
        [''],
        ['统计项目', '数值', '单位'],
        ['巡检记录数', summary_stats['records'] or 0, '条'],
        ['涉及结晶池', summary_stats['pools'] or 0, '个'],
        ['批次号数', summary_stats['batches'] or 0, '个'],
        ['平均卤水浓度', round(float(summary_stats['avg_conc'] or 0), 2), '°Bé'],
        ['最低卤水浓度', round(float(summary_stats['min_conc'] or 0), 2), '°Bé'],
        ['最高卤水浓度', round(float(summary_stats['max_conc'] or 0), 2), '°Bé'],
        ['累计收盐量', round(float(summary_stats['total_yield'] or 0), 3), '吨'],
        ['异常记录数', summary_stats['anomalies'] or 0, '条'],
    ]
    for r_idx, row in enumerate(summary_data, 1):
        for c_idx, val in enumerate(row, 1):
            c = summary_ws.cell(row=r_idx, column=c_idx, value=val)
            if r_idx == 1:
                c.font = Font(bold=True, size=16)
            elif r_idx == 5:
                c.font = Font(bold=True)
                c.fill = header_fill
                c.font = Font(bold=True, color='FFFFFF')
            c.border = border

    summary_ws.column_dimensions['A'].width = 22
    summary_ws.column_dimensions['B'].width = 22
    summary_ws.column_dimensions['C'].width = 12
    summary_ws.merge_cells('A1:C1')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="巡检台账_{today_str}.xlsx"'
    wb.save(response)
    return response


def export_summary(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        messages.error(request, '请先安装 openpyxl：pip install openpyxl')
        return redirect('inspection:record_list')

    today = timezone.localdate()
    last_30 = today - timedelta(days=30)

    wb = Workbook()
    ws = wb.active
    ws.title = '综合摘要'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
    sub_header_fill = PatternFill(start_color='99F6E4', end_color='99F6E4', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(border_style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    ws.merge_cells('A1:F1')
    title = ws.cell(row=1, column=1, value='海盐晒场结晶池巡检台账 - 综合摘要报告')
    title.font = Font(bold=True, size=18)
    title.alignment = center

    info = [
        ['报告生成时间', timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')],
        ['统计区间', f'{last_30.isoformat()} ~ {today.isoformat()}'],
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=v)

    r = 6
    ws.merge_cells(f'A{r}:F{r}')
    ws.cell(row=r, column=1, value='一、总体概况').font = Font(bold=True, size=14)

    total_stats = InspectionRecord.objects.filter(inspection_date__gte=last_30).aggregate(
        records=Count('id'),
        pools=Count('pool', distinct=True),
        avg_conc=Avg('brine_concentration'),
        min_conc=Min('brine_concentration'),
        max_conc=Max('brine_concentration'),
        total_yield=Sum('salt_yield'),
        anomalies=Count('id', filter=Q(has_anomaly=True)),
    )
    pool_active = CrystallizationPool.objects.filter(status='active').count()
    alert_open = AnomalyAlert.objects.filter(status__in=('open', 'processing')).count()

    r = 8
    overview = [
        ['指标', '数值', '单位', '', '', ''],
        ['结晶池总数', CrystallizationPool.objects.count(), '个', '生产中', pool_active, '个'],
        ['30天巡检记录', total_stats['records'] or 0, '条', '异常记录', total_stats['anomalies'] or 0, '条'],
        ['平均卤水浓度', round(float(total_stats['avg_conc'] or 0), 2), '°Bé', '待处理告警', alert_open, '条'],
        ['累计收盐量(30天)', round(float(total_stats['total_yield'] or 0), 3), '吨', '浓度范围',
         f"{round(float(total_stats['min_conc'] or 0), 2)}~{round(float(total_stats['max_conc'] or 0), 2)}", '°Bé'],
    ]
    for i, row in enumerate(overview):
        for j, val in enumerate(row):
            c = ws.cell(row=r + i, column=j + 1, value=val)
            c.border = border
            c.alignment = center
            if i == 0:
                c.fill = header_fill
                c.font = header_font

    r += len(overview) + 2
    ws.merge_cells(f'A{r}:F{r}')
    ws.cell(row=r, column=1, value='二、分池组统计').font = Font(bold=True, size=14)
    r += 2
    group_headers = ['池组', '结晶池数', '生产中', '巡检次数', '平均浓度(°Bé)', '累计收盐(吨)']
    for i, h in enumerate(group_headers):
        c = ws.cell(row=r, column=i + 1, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = center

    groups = ['A', 'B', 'C', 'D']
    for gi, g in enumerate(groups, r + 1):
        g_pools = CrystallizationPool.objects.filter(pool_group=g)
        g_ids = g_pools.values_list('id', flat=True)
        g_records = InspectionRecord.objects.filter(pool_id__in=g_ids, inspection_date__gte=last_30)
        gs = g_records.aggregate(avg=Avg('brine_concentration'), y=Sum('salt_yield'))
        row_vals = [g, g_pools.count(), g_pools.filter(status='active').count(),
                    g_records.count(), round(float(gs['avg'] or 0), 2), round(float(gs['y'] or 0), 3)]
        for j, val in enumerate(row_vals):
            c = ws.cell(row=gi, column=j + 1, value=val)
            c.border = border
            c.alignment = center

    r = gi + 3
    ws.merge_cells(f'A{r}:F{r}')
    ws.cell(row=r, column=1, value='三、结晶池状态明细').font = Font(bold=True, size=14)
    r += 2
    pool_headers = ['池号', '池名', '状态', '最新浓度(°Bé)', '累计产量(吨)', '最近巡检']
    for i, h in enumerate(pool_headers):
        c = ws.cell(row=r, column=i + 1, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = center

    pools = CrystallizationPool.objects.all().order_by('pool_code')
    red_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    for pi, p in enumerate(pools, r + 1):
        latest = p.inspections.first()
        y_sum = p.inspections.filter(inspection_date__gte=last_30).aggregate(s=Sum('salt_yield'))['s'] or 0
        row_vals = [
            p.pool_code, p.pool_name, p.get_status_display(),
            float(latest.brine_concentration) if latest else '-',
            round(float(y_sum), 3),
            latest.inspection_date.isoformat() if latest else '-',
        ]
        is_anomaly = latest and latest.has_anomaly
        for j, val in enumerate(row_vals):
            c = ws.cell(row=pi, column=j + 1, value=val)
            c.border = border
            c.alignment = center
            if is_anomaly:
                c.fill = red_fill

    widths = [16, 16, 14, 16, 16, 16]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w

    alert_ws = wb.create_sheet('异常清单')
    alert_headers = ['告警ID', '池号', '类型', '严重程度', '标题', '状态', '发现时间']
    for i, h in enumerate(alert_headers):
        c = alert_ws.cell(row=1, column=i + 1, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = center
    for ai, a in enumerate(AnomalyAlert.objects.select_related('pool').order_by('-created_at')[:500], 2):
        vals = [a.id, a.pool.pool_code, a.get_anomaly_type_display(), a.get_severity_display(),
                a.title, a.get_status_display(), a.detected_at.strftime('%Y-%m-%d %H:%M')]
        for j, val in enumerate(vals):
            c = alert_ws.cell(row=ai, column=j + 1, value=val)
            c.border = border
            c.alignment = center
    for col, w in enumerate([10, 12, 12, 12, 40, 12, 20], 1):
        alert_ws.column_dimensions[chr(64 + col)].width = w

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="巡检摘要报告_{today.isoformat()}.xlsx"'
    wb.save(response)
    return response
