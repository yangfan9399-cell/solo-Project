from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db.models import Q, Sum
from django.utils import timezone
from datetime import date, timedelta
import csv
import io

from .models import (
    CoreSample, Cutter, CuttingPurpose, CuttingTask,
    BatchVersion, AnomalyRecord
)
from .services import (
    detect_all_anomalies, get_daily_schedule, generate_summary_data,
    create_batch_version
)
from .forms import (
    CoreSampleForm, CuttingTaskForm, BatchVersionForm,
    AnomalyResolveForm
)


def dashboard(request):
    detect_all_anomalies()
    summary = generate_summary_data()
    recent_tasks = CuttingTask.objects.select_related(
        'core_sample', 'cutter', 'purpose'
    ).order_by('-created_at')[:10]
    open_anomalies = AnomalyRecord.objects.filter(
        resolved=False
    ).select_related('core_sample', 'cutting_task', 'cutter')[:5]
    today_schedule = get_daily_schedule(date.today())

    context = {
        'summary': summary,
        'recent_tasks': recent_tasks,
        'open_anomalies': open_anomalies,
        'today_schedule': today_schedule,
        'today': date.today(),
    }
    return render(request, 'core/dashboard.html', context)


def workbench(request):
    target_date_str = request.GET.get('date', '')
    if target_date_str:
        try:
            target_date = date.fromisoformat(target_date_str)
        except ValueError:
            target_date = date.today()
    else:
        target_date = date.today()

    detect_all_anomalies()
    schedule = get_daily_schedule(target_date)
    pending_tasks = CuttingTask.objects.filter(
        status__in=['pending', 'scheduled'],
        scheduled_date__isnull=True
    ).select_related('core_sample', 'purpose').order_by(
        '-core_sample__priority', 'created_at'
    )[:15]

    context = {
        'schedule': schedule,
        'pending_tasks': pending_tasks,
        'target_date': target_date,
        'prev_date': target_date - timedelta(days=1),
        'next_date': target_date + timedelta(days=1),
    }
    return render(request, 'core/workbench.html', context)


def sample_list(request):
    query = request.GET.get('q', '')
    priority = request.GET.get('priority', '')
    status = request.GET.get('status', '')
    lithology = request.GET.get('lithology', '')

    samples = CoreSample.objects.prefetch_related('cutting_tasks', 'anomalies')

    if query:
        samples = samples.filter(
            Q(sample_no__icontains=query) |
            Q(well_name__icontains=query) |
            Q(description__icontains=query)
        )
    if priority:
        samples = samples.filter(priority=priority)
    if status:
        samples = samples.filter(status=status)
    if lithology:
        samples = samples.filter(lithology__icontains=lithology)

    samples = samples.order_by('-priority', 'sample_no')

    lithologies = CoreSample.objects.values_list('lithology', flat=True).distinct().exclude(
        lithology=''
    ).order_by('lithology')

    context = {
        'samples': samples,
        'query': query,
        'priority': priority,
        'status': status,
        'lithology': lithology,
        'lithologies': lithologies,
        'priority_choices': CoreSample.PRIORITY_CHOICES,
        'status_choices': CoreSample.STATUS_CHOICES,
    }
    return render(request, 'core/sample_list.html', context)


def sample_detail(request, pk):
    sample = get_object_or_404(CoreSample, pk=pk)
    tasks = sample.cutting_tasks.select_related('cutter', 'purpose', 'batch_version').order_by('-created_at')
    anomalies = sample.anomalies.all()[:10]
    context = {
        'sample': sample,
        'tasks': tasks,
        'anomalies': anomalies,
    }
    return render(request, 'core/sample_detail.html', context)


def sample_create(request):
    if request.method == 'POST':
        form = CoreSampleForm(request.POST)
        if form.is_valid():
            sample = form.save()
            messages.success(request, f'岩心样本 {sample.sample_no} 创建成功')
            return redirect('core:sample_detail', pk=sample.pk)
    else:
        form = CoreSampleForm()
    context = {'form': form, 'title': '新建岩心样本'}
    return render(request, 'core/sample_form.html', context)


def sample_edit(request, pk):
    sample = get_object_or_404(CoreSample, pk=pk)
    if request.method == 'POST':
        form = CoreSampleForm(request.POST, instance=sample)
        if form.is_valid():
            form.save()
            messages.success(request, '岩心样本更新成功')
            return redirect('core:sample_detail', pk=sample.pk)
    else:
        form = CoreSampleForm(instance=sample)
    context = {'form': form, 'sample': sample, 'title': '编辑岩心样本'}
    return render(request, 'core/sample_form.html', context)


def task_list(request):
    query = request.GET.get('q', '')
    status = request.GET.get('status', '')
    purpose = request.GET.get('purpose', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')

    tasks = CuttingTask.objects.select_related('core_sample', 'cutter', 'purpose', 'batch_version')

    if query:
        tasks = tasks.filter(
            Q(task_no__icontains=query) |
            Q(core_sample__sample_no__icontains=query) |
            Q(remarks__icontains=query)
        )
    if status:
        tasks = tasks.filter(status=status)
    if purpose:
        tasks = tasks.filter(purpose_id=purpose)
    if date_from:
        tasks = tasks.filter(scheduled_date__gte=date_from)
    if date_to:
        tasks = tasks.filter(scheduled_date__lte=date_to)

    tasks = tasks.order_by('-scheduled_date', '-scheduled_start_time')

    purposes = CuttingPurpose.objects.filter(is_active=True)

    context = {
        'tasks': tasks,
        'query': query,
        'status': status,
        'purpose': purpose,
        'date_from': date_from,
        'date_to': date_to,
        'purposes': purposes,
        'status_choices': CuttingTask.STATUS_CHOICES,
    }
    return render(request, 'core/task_list.html', context)


def task_detail(request, pk):
    task = get_object_or_404(CuttingTask, pk=pk)
    anomalies = task.anomalies.all()
    context = {
        'task': task,
        'anomalies': anomalies,
    }
    return render(request, 'core/task_detail.html', context)


def task_create(request):
    if request.method == 'POST':
        form = CuttingTaskForm(request.POST)
        if form.is_valid():
            task = form.save()
            messages.success(request, f'切割任务 {task.task_no} 创建成功')
            return redirect('core:task_detail', pk=task.pk)
    else:
        form = CuttingTaskForm()
    context = {'form': form, 'title': '新建切割任务'}
    return render(request, 'core/task_form.html', context)


def task_edit(request, pk):
    task = get_object_or_404(CuttingTask, pk=pk)
    if request.method == 'POST':
        form = CuttingTaskForm(request.POST, instance=task)
        if form.is_valid():
            form.save()
            detect_all_anomalies()
            messages.success(request, '切割任务更新成功')
            return redirect('core:task_detail', pk=task.pk)
    else:
        form = CuttingTaskForm(instance=task)
    context = {'form': form, 'task': task, 'title': '编辑切割任务'}
    return render(request, 'core/task_form.html', context)


def batch_list(request):
    batches = BatchVersion.objects.prefetch_related('tasks').order_by('-created_at')
    context = {'batches': batches}
    return render(request, 'core/batch_list.html', context)


def batch_detail(request, pk):
    batch = get_object_or_404(BatchVersion, pk=pk)
    tasks = batch.tasks.select_related('core_sample', 'cutter', 'purpose').order_by('scheduled_start_time')
    snapshot = batch.get_snapshot()
    prev_batches = BatchVersion.objects.filter(
        batch_no=batch.batch_no,
        created_at__lt=batch.created_at
    ).order_by('-created_at')[:5]

    context = {
        'batch': batch,
        'tasks': tasks,
        'snapshot': snapshot,
        'prev_batches': prev_batches,
    }
    return render(request, 'core/batch_detail.html', context)


def batch_create(request):
    if request.method == 'POST':
        form = BatchVersionForm(request.POST)
        if form.is_valid():
            batch = create_batch_version(
                batch_no=form.cleaned_data['batch_no'],
                batch_type=form.cleaned_data['batch_type'],
                description=form.cleaned_data['description'],
                created_by=form.cleaned_data.get('created_by', ''),
            )
            messages.success(request, f'批次 {batch.version_no} 创建成功')
            return redirect('core:batch_detail', pk=batch.pk)
    else:
        form = BatchVersionForm()
    context = {'form': form, 'title': '创建批次版本'}
    return render(request, 'core/batch_form.html', context)


def anomaly_list(request):
    severity = request.GET.get('severity', '')
    anomaly_type = request.GET.get('type', '')
    resolved = request.GET.get('resolved', '')

    anomalies = AnomalyRecord.objects.select_related(
        'core_sample', 'cutting_task', 'cutter'
    )

    if severity:
        anomalies = anomalies.filter(severity=severity)
    if anomaly_type:
        anomalies = anomalies.filter(anomaly_type=anomaly_type)
    if resolved == 'yes':
        anomalies = anomalies.filter(resolved=True)
    elif resolved == 'no':
        anomalies = anomalies.filter(resolved=False)

    anomalies = anomalies.order_by('-severity', '-detected_at')

    context = {
        'anomalies': anomalies,
        'severity': severity,
        'anomaly_type': anomaly_type,
        'resolved': resolved,
        'severity_choices': AnomalyRecord.SEVERITY_CHOICES,
        'type_choices': AnomalyRecord.TYPE_CHOICES,
    }
    return render(request, 'core/anomaly_list.html', context)


def anomaly_resolve(request, pk):
    anomaly = get_object_or_404(AnomalyRecord, pk=pk)
    if request.method == 'POST':
        form = AnomalyResolveForm(request.POST)
        if form.is_valid():
            anomaly.resolved = True
            anomaly.resolution = form.cleaned_data['resolution']
            anomaly.save()
            messages.success(request, '异常已标记为已解决')
            return redirect('core:anomaly_list')
    else:
        form = AnomalyResolveForm()
    context = {'form': form, 'anomaly': anomaly}
    return render(request, 'core/anomaly_resolve.html', context)


def export_summary(request):
    detect_all_anomalies()
    summary = generate_summary_data()

    today = date.today()
    tasks = CuttingTask.objects.filter(
        scheduled_date=today
    ).select_related('core_sample', 'cutter', 'purpose').order_by('scheduled_start_time')

    samples = CoreSample.objects.all()

    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="core_scheduler_summary_{today}.csv"'

    response.write('\ufeff')

    writer = csv.writer(response)
    writer.writerow(['岩心样本切割排程系统 - 摘要报告'])
    writer.writerow(['生成日期', today.isoformat()])
    writer.writerow([])

    writer.writerow(['一、总体概览'])
    writer.writerow(['岩心样本总数', summary['total_samples']])
    writer.writerow(['待切割', summary['pending_samples']])
    writer.writerow(['切割中', summary['in_progress_samples']])
    writer.writerow(['已完成', summary['completed_samples']])
    writer.writerow(['切割任务总数', summary['total_tasks']])
    writer.writerow(['已完成任务', summary['completed_tasks']])
    writer.writerow(['累计切割长度(m)', summary['total_cut_length']])
    writer.writerow(['累计损耗(m)', summary['total_loss']])
    writer.writerow(['未处理异常数', summary['open_anomalies']])
    writer.writerow([])

    writer.writerow(['二、今日排程'])
    writer.writerow(['任务编号', '岩心样本', '切割机', '切割目的', '计划长度(m)', '状态', '开始时间', '结束时间'])
    for task in tasks:
        writer.writerow([
            task.task_no,
            task.core_sample.sample_no if task.core_sample else '',
            task.cutter.cutter_no if task.cutter else '',
            task.purpose.name if task.purpose else '',
            task.planned_cut_length,
            task.get_status_display(),
            task.scheduled_start_time or '',
            task.scheduled_end_time or '',
        ])
    writer.writerow([])

    writer.writerow(['三、岩心样本台账'])
    writer.writerow(['样本编号', '井号', '总长度(m)', '剩余长度(m)', '使用率(%)', '优先级', '状态', '岩性', '地层'])
    for sample in samples:
        writer.writerow([
            sample.sample_no,
            sample.well_name,
            sample.total_length,
            sample.remaining_length,
            sample.usage_rate,
            sample.get_priority_display(),
            sample.get_status_display(),
            sample.lithology,
            sample.formation,
        ])

    return response


def export_summary_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    detect_all_anomalies()
    summary = generate_summary_data()
    today = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = '摘要报告'

    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    ws['A1'] = '岩心样本切割排程系统 - 摘要报告'
    ws['A1'].font = title_font
    ws.merge_cells('A1:F1')
    ws['A2'] = f'生成日期: {today.isoformat()}'
    ws.merge_cells('A2:F2')

    row = 4
    ws.cell(row=row, column=1, value='一、总体概览').font = Font(bold=True, size=12)
    row += 1
    overview_data = [
        ('岩心样本总数', summary['total_samples']),
        ('待切割', summary['pending_samples']),
        ('切割中', summary['in_progress_samples']),
        ('已完成', summary['completed_samples']),
        ('切割任务总数', summary['total_tasks']),
        ('已完成任务', summary['completed_tasks']),
        ('累计切割长度(m)', summary['total_cut_length']),
        ('累计损耗(m)', summary['total_loss']),
        ('未处理异常数', summary['open_anomalies']),
        ('切割机数量', summary['cutters']),
        ('可用切割机', summary['available_cutters']),
    ]
    for label, value in overview_data:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value='二、今日排程').font = Font(bold=True, size=12)
    row += 1
    schedule_headers = ['任务编号', '岩心样本', '切割机', '切割目的', '计划长度(m)', '状态', '开始时间', '结束时间']
    for col, header in enumerate(schedule_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    row += 1

    tasks = CuttingTask.objects.filter(
        scheduled_date=today
    ).select_related('core_sample', 'cutter', 'purpose').order_by('scheduled_start_time')
    for task in tasks:
        ws.cell(row=row, column=1, value=task.task_no).border = thin_border
        ws.cell(row=row, column=2, value=task.core_sample.sample_no if task.core_sample else '').border = thin_border
        ws.cell(row=row, column=3, value=task.cutter.cutter_no if task.cutter else '').border = thin_border
        ws.cell(row=row, column=4, value=task.purpose.name if task.purpose else '').border = thin_border
        ws.cell(row=row, column=5, value=task.planned_cut_length).border = thin_border
        ws.cell(row=row, column=6, value=task.get_status_display()).border = thin_border
        ws.cell(row=row, column=7, value=str(task.scheduled_start_time) if task.scheduled_start_time else '').border = thin_border
        ws.cell(row=row, column=8, value=str(task.scheduled_end_time) if task.scheduled_end_time else '').border = thin_border
        row += 1

    row += 1
    ws.cell(row=row, column=1, value='三、岩心样本台账').font = Font(bold=True, size=12)
    row += 1
    sample_headers = ['样本编号', '井号', '总长度(m)', '剩余长度(m)', '使用率(%)', '优先级', '状态', '岩性', '地层']
    for col, header in enumerate(sample_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    row += 1

    samples = CoreSample.objects.all().order_by('sample_no')
    for sample in samples:
        ws.cell(row=row, column=1, value=sample.sample_no).border = thin_border
        ws.cell(row=row, column=2, value=sample.well_name).border = thin_border
        ws.cell(row=row, column=3, value=sample.total_length).border = thin_border
        ws.cell(row=row, column=4, value=sample.remaining_length).border = thin_border
        ws.cell(row=row, column=5, value=sample.usage_rate).border = thin_border
        ws.cell(row=row, column=6, value=sample.get_priority_display()).border = thin_border
        ws.cell(row=row, column=7, value=sample.get_status_display()).border = thin_border
        ws.cell(row=row, column=8, value=sample.lithology).border = thin_border
        ws.cell(row=row, column=9, value=sample.formation).border = thin_border
        row += 1

    for col in range(1, 10):
        ws.column_dimensions[chr(64 + col)].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="core_scheduler_summary_{today}.xlsx"'
    return response


def cutter_list(request):
    cutters = Cutter.objects.all().order_by('cutter_no')
    context = {'cutters': cutters}
    return render(request, 'core/cutter_list.html', context)


def purpose_list(request):
    purposes = CuttingPurpose.objects.all().order_by('sort_order', 'code')
    context = {'purposes': purposes}
    return render(request, 'core/purpose_list.html', context)
